
from pathlib import Path
from docx import Document


class ReportGenerator:
    VERSION = "report-generator-v1"

    DISCLAIMER = (
        "This report is generated for legal intelligence and research support. "
        "It is not legal advice, does not create an attorney-client relationship, "
        "and must be reviewed by a qualified legal professional before reliance "
        "or action."
    )

    def __init__(self, output_dir):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, state, filename=None):
        if not isinstance(state, dict):
            raise TypeError("state must be a dictionary")

        document = Document()

        document.add_heading("LegalMind Legal Intelligence Report", level=0)

        document.add_heading("Report Metadata", level=1)
        self._add_field(document, "Generator", self.VERSION)
        self._add_field(
            document,
            "Document ID",
            state.get("document_id"),
        )
        self._add_field(
            document,
            "Workflow Status",
            state.get("workflow_status"),
        )

        metadata = state.get("metadata") or {}
        self._add_field(
            document,
            "Corpus",
            metadata.get("corpus")
            or state.get("corpus"),
        )
        self._add_field(
            document,
            "Jurisdiction",
            metadata.get("jurisdiction")
            or state.get("jurisdiction"),
        )

        document.add_heading("Risk Findings", level=1)
        risk_findings = state.get("risk_findings") or []
        if risk_findings:
            for finding in risk_findings:
                document.add_paragraph(str(finding))
        else:
            document.add_paragraph("No risk findings available.")

        document.add_heading("Legal Reasoning", level=1)
        reasoning = state.get("reasoning_result")
        document.add_paragraph(
            str(reasoning)
            if reasoning is not None
            else "No legal reasoning result available."
        )

        document.add_heading("Critic Review", level=1)
        critic = state.get("critic_result")
        document.add_paragraph(
            str(critic)
            if critic is not None
            else "No critic result available."
        )

        document.add_heading("Explainability", level=1)
        explainability = state.get("explainability_result")
        document.add_paragraph(
            str(explainability)
            if explainability is not None
            else "No explainability result available."
        )

        document.add_heading("Evidence and Traceability", level=1)
        research = state.get("research_result") or {}
        results = research.get("results") or []

        evidence_count = 0
        case_count = 0

        for item in results:
            result = item.get("result") or {}
            evidence = result.get("evidence") or []
            case_law = result.get("ranked_case_law") or []

            evidence_count += len(evidence)
            case_count += len(case_law)

            query = result.get("query") or item.get("clause", {}).get("text")
            if query:
                document.add_paragraph(
                    f"Query: {query}",
                    style="List Bullet",
                )

            for evidence_item in evidence:
                document.add_paragraph(
                    str(evidence_item),
                    style="List Bullet 2",
                )

            for case in case_law:
                document.add_paragraph(
                    f"Case-law evidence: {case}",
                    style="List Bullet 2",
                )

        self._add_field(document, "Evidence items", evidence_count)
        self._add_field(document, "Case-law items", case_count)

        document.add_heading("Human Review", level=1)
        hitl = state.get("hitl_result")
        document.add_paragraph(
            str(hitl)
            if hitl is not None
            else "No HITL review result recorded."
        )

        document.add_heading("Disclaimer", level=1)
        document.add_paragraph(self.DISCLAIMER)

        if filename is None:
            filename = f"{state.get('document_id', 'legalmind_report')}.docx"

        output_path = self.output_dir / filename
        document.save(output_path)

        return {
            "status": "COMPLETED",
            "format": "DOCX",
            "path": str(output_path),
            "generator": self.VERSION,
            "document_id": state.get("document_id"),
            "corpus": metadata.get("corpus") or state.get("corpus"),
            "jurisdiction": (
                metadata.get("jurisdiction")
                or state.get("jurisdiction")
            ),
            "evidence_count": evidence_count,
            "case_law_count": case_count,
            "disclaimer_included": True,
        }

    @staticmethod
    def _add_field(document, label, value):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value) if value is not None else "Not available")
